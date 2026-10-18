"""Editable native PowerPoint artifacts + real LibreOffice PDF exports + technical PDF.
All metrics/screenshots are read from executed project outputs; no invented chart data.
"""
import json, re, subprocess
from pathlib import Path
from copy import deepcopy
from html import escape
from PIL import Image
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.chart.data import CategoryChartData,XyChartData
from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION,XL_LABEL_POSITION
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak,KeepTogether
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import pymupdf as fitz
from backend.app.config import ROOT,ARTIFACTS,DATA
DOCS=ROOT/'docs';PDIR=DOCS/'presentation';ADIR=DOCS/'architecture';ASSETS=PDIR/'assets'
BG='F5F8F4';INK='163E35';TEAL='287E68';MUTED='6C8274';LINE='DAE4D9';LIME='CEE2A9';WHITE='FFFFFF';AMBER='AC8552';CORAL='B87561'
FONT='Liberation Sans'

def rgb(v):return RGBColor.from_string(v)
def rect(slide,x,y,w,h,fill=WHITE,line=None,radius=True):
    s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid();s.fill.fore_color.rgb=rgb(fill)
    if line:s.line.color.rgb=rgb(line);s.line.width=Pt(.7)
    else:s.line.fill.background()
    if radius:
        try:s.adjustments[0]=.12
        except Exception:pass
    return s

def text(slide,x,y,w,h,value,size=20,color=INK,bold=False,align=PP_ALIGN.LEFT):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_top=0;tf.margin_bottom=0;tf.margin_left=0;tf.margin_right=0
    for i,line in enumerate(str(value).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.alignment=align;p.space_after=Pt(5);p.line_spacing=1.07
    return box

def label(slide,x,y,w,value,dark=False):
    rect(slide,x,y,w,.32,fill='294F41' if dark else 'E7F0E1',radius=True)
    text(slide,x+.11,y+.065,w-.2,.18,value,size=8.5,color=LIME if dark else TEAL,bold=True)

def footer(slide,n,total=16,dark=False):
    col='9BB6A4' if dark else '87977F'
    text(slide,.55,7.13,10.5,.18,'PRANAGUARD   /   SYNTHETIC PROOF OF CONCEPT · NOT CLINICALLY VALIDATED',8,color=col)
    text(slide,12.1,7.11,.65,.2,f'{n:02d} / {total:02d}',8,color=col,align=PP_ALIGN.RIGHT)

def new_slide(prs,title,kicker,n,subtitle=None,dark=False):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(INK if dark else BG)
    text(s,.55,.35,11,.2,kicker.upper(),9,color=LIME if dark else TEAL,bold=True)
    text(s,.55,.82,12.2,.78,title,30,color=WHITE if dark else INK,bold=True)
    if subtitle:text(s,.57,1.58,12.1,.48,subtitle,13,color='BACFBC' if dark else MUTED)
    footer(s,n,dark=dark)
    s.notes_slide.notes_text_frame.text='PranaGuard research prototype. All patient physiology/outcomes synthetic. No clinical validation. Use docs/demo/DEMO_SCRIPT.md for the timed narration. Actual metrics: artifacts/evaluation.json. Environmental attribution: Open-Meteo / CAMS, CC BY 4.0. Public repository/video status remains pending until actually verified.'
    return s

def bullets(slide,x,y,w,items,size=19,spacing=.84,color=INK):
    for i,item in enumerate(items):
        rect(slide,x,y+i*spacing+.1,.065,.065,TEAL,radius=False)
        text(slide,x+.24,y+i*spacing,w-.24,spacing-.08,item,size,color=color)

def callout(slide,x,y,w,h,title,body,fill='E9F1E2',accent=TEAL):
    rect(slide,x,y,w,h,fill=fill,line=LINE)
    text(slide,x+.22,y+.18,w-.44,.42,title,17,color=accent,bold=True)
    text(slide,x+.22,y+.75,w-.44,h-.9,body,14,color=MUTED)

def add_table(slide,x,y,w,h,headers,rows,font=14,widths=None):
    sh=slide.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(w),Inches(h));tb=sh.table
    if widths:
        for col,ww in zip(tb.columns,widths):col.width=Inches(ww)
    for r,values in enumerate([headers]+rows):
        for c,value in enumerate(values):
            cell=tb.cell(r,c);cell.text=str(value);cell.margin_top=Inches(.06);cell.margin_bottom=Inches(.04);cell.margin_left=Inches(.12);cell.margin_right=Inches(.08);cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid();cell.fill.fore_color.rgb=rgb(INK if r==0 else WHITE if r%2 else 'EDF3E8')
            for p in cell.text_frame.paragraphs:
                p.font.name=FONT;p.font.size=Pt(font if r else max(10,font-1));p.font.bold=r==0;p.font.color.rgb=rgb(WHITE if r==0 else INK)
    return tb

def picture_fit(slide,path,x,y,w,h):
    im=Image.open(path);ratio=im.width/im.height;outer=w/h
    if ratio>outer:pw=w;ph=w/ratio
    else:ph=h;pw=h*ratio
    rect(slide,x-.025,y-.025,w+.05,h+.05,WHITE,LINE)
    return slide.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))

def crop(name,box,target):
    im=Image.open(DOCS/'screenshots'/name);im.crop(box).save(ASSETS/target);return ASSETS/target

def connector(slide,x1,y1,x2,y2,color='799885',dashed=False,arrow=True):
    line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));line.line.color.rgb=rgb(color);line.line.width=Pt(1.4 if not dashed else .9)
    if dashed:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        line.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    from lxml import etree
    ln=line._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln')
    if arrow:
        tail=etree.SubElement(ln,'{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd');tail.set('type','triangle');tail.set('w','sm');tail.set('len','sm')
    return line

ARCH_NODES=[
 ('01','Data sources','Synthetic clinical/body/symptom streams\nReal cached ambient/weather fields','generate_data.py · environment cache'),
 ('02','Ingestion & replay','Seeded hourly observations · UTC clock\nDeterministic cohort-wide updates','engine.py · replay task'),
 ('03','Validation & features','Schema/ranges/time alignment\nCausal 6h/24h windows; train-only imputation','validation.py · features.py'),
 ('04','Digital-twin state','Persistent patient identity + baseline\nCurrent state, missingness and versions','engine.py · first72h baseline'),
 ('05','ML & calibration','Separate 24h / 48h / 72h LightGBM\nSeparate Platt models; LR baseline','ml.py · models.joblib'),
 ('06','Explainability & alerts','On-demand native TreeSHAP\nThreshold alerts and contribution deltas','ml.py · engine.py'),
 ('07','What-if simulation','Clone observed past; projected branches\nNo future rows; original state preserved','engine.py · scenarios table'),
 ('08','SQLite persistence','States + cursor + alerts + actions\nSeparate scenarios; atomic transactions','store.py · SQLite WAL'),
 ('09','FastAPI APIs / WS','REST inference/review/replay endpoints\nWebSocket invalidation; optional write token','main.py · /docs · /api/ws'),
 ('10','React dashboard','7 clinician views · actual API charts\nSearch, review, replay, scenarios, export','React · TS · Vite · Tailwind · Recharts')]

def architecture(slide,scale=1,ox=0,oy=0,header=True):
    def tx(x,y,w,h,val,size=14,col=INK,bold=False):return text(slide,ox+x*scale,oy+y*scale,w*scale,h*scale,val,max(9,size*scale) if scale<1 else size,col,bold)
    def bx(x,y,w,h,fill=WHITE,line=LINE):return rect(slide,ox+x*scale,oy+y*scale,w*scale,h*scale,fill,line)
    if header:
        tx(.55,.36,14,.25,'PRANAGUARD  /  IMPLEMENTED SYSTEM ARCHITECTURE',11,TEAL,True)
        tx(.55,.8,14.6,.6,'A persistent twin, not a static score',30,INK,True)
    coords=[(.55,1.8),(4.37,1.8),(8.19,1.8),(12.01,1.8),(12.01,3.85),(8.19,3.85),(4.37,3.85),(.55,3.85),(.55,5.9),(4.37,5.9)]
    w=3.4;h=1.52
    # Arrows underneath boxes. Logical chain plus truthful cross-cutting runtime branches.
    for i in range(9):
        a=coords[i];b=coords[i+1]
        if a[1]==b[1]:
            right=b[0]>a[0];x1=a[0]+(w if right else 0);x2=b[0]+(0 if right else w);y1=y2=a[1]+h/2
        else:x1=x2=a[0]+w/2;y1=a[1]+h;y2=b[1]
        connector(slide,ox+x1*scale,oy+y1*scale,ox+x2*scale,oy+y2*scale)
    # Direct prediction persistence, state-to-scenario and API-triggered calculations are not a strictly serial chain.
    for a,b,end in [((13.71,5.37),(13.71,5.6),False),((13.71,5.6),(2.25,5.6),False),((2.25,5.6),(2.25,5.37),True)]:
        connector(slide,ox+a[0]*scale,oy+a[1]*scale,ox+b[0]*scale,oy+b[1]*scale,dashed=True,arrow=end)
    tx(8.5,5.36,4,.18,'prediction → atomic state persistence',9,MUTED)
    for a,b,end in [((13.71,1.8),(13.71,1.58),False),((13.71,1.58),(7.98,1.58),False),((7.98,1.58),(7.98,3.6),False),((7.98,3.6),(6.07,3.6),False),((6.07,3.6),(6.07,3.85),True)]:
        connector(slide,ox+a[0]*scale,oy+a[1]*scale,ox+b[0]*scale,oy+b[1]*scale,dashed=True,arrow=end)
    tx(8.05,1.37,5,.18,'twin observed past → independent scenario branch',9,MUTED)
    compact=[('Synthetic body/clinical streams\nCached real ambient/weather','Synthetic + cached inputs'),('Hourly cohort-wide replay\nSeeded UTC clock','engine.py'),('Ranges, duplicates, timestamps\nCausal6h /24h windows','validation.py · features.py'),('Identity + first72h baseline\nCurrent data + version history','engine.py'),('Three LightGBM + Platt models\n24h /48h /72h · LR baseline','ml.py · models.joblib'),('Native TreeSHAP on demand\nPersisted threshold alerts','ml.py · engine.py'),('Cloned past + projected inputs\nSeparate branches, no future data','engine.py'),('States, cursor, alerts, actions\nScenarios saved separately','store.py · SQLite WAL'),('REST + WebSocket invalidation\nOptional write-access token','main.py · /docs'),('7 working clinician views\nReal API charts and controls','React clinician frontend')]
    for j,((idx,title,body,module),(x,y)) in enumerate(zip(ARCH_NODES,coords)):
        if scale<1:body,module=compact[j]
        bx(x,y,w,h,fill=WHITE if int(idx)%2 else 'EBF2E6')
        tx(x+.14,y+.12,.37,.26,idx,10,TEAL,True)
        tx(x+.57,y+.10,w-.7,.35,title,16,INK,True)
        tx(x+.16,y+.59,w-.32,.54,body,12 if scale<1 else 10.2,MUTED)
        tx(x+.16,y+1.25,w-.32,.17,module,8.1,TEAL)
    bx(8.19,5.9,7.22,1.52,fill='E9F0E2')
    tx(8.39,6.1,6.8,.3,'Offline training & evaluation artifacts',17,INK,True)
    tx(8.39,6.65,6.8,.55,'15 executed experiments · patient/time-disjoint splits · ≥72h purge\nModels, calibration, ROC/PR, ablations and provenance versioned',11,MUTED)
    tx(.55,7.88,14.9,.65,'Solid arrows: logical capability chain. Dashed: cross-cutting runtime branches. SQLite writes occur on every twin update; explanations and scenarios are API-triggered. All patients synthetic; cached ambient data are historical. The 72h model is weak and its calibration slope is negative.',11,MUTED)
    tx(.55,8.65,14.9,.18,'RESEARCH PROTOTYPE · NO CLINICAL VALIDATION · ORIGINAL EDITABLE VECTOR SHAPES',9,TEAL,True)

def build_architecture():
    prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9);s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG);architecture(s)
    prs.save(ADIR/'PranaGuard_Architecture.pptx')

def build_deck():
    r=json.loads((ARTIFACTS/'evaluation.json').read_text());g=json.loads((DATA/'generation_manifest.json').read_text());meta=json.loads((ROOT/'submission_metadata.json').read_text())
    team=meta['team_name'] or '[TEAM NAME PENDING]';institution=meta['college_or_incubator'] or '[COLLEGE / INCUBATOR PENDING]'
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    prs.core_properties.title='PranaGuard — Body × Environment Digital Twin';prs.core_properties.subject='Synthetic COPD research prototype; actual negative and positive experiment results';prs.core_properties.author='PranaGuard contributors (team details pending)'
    overview=crop('01_overview.png',(260,205,1490,840),'overview_crop.png')
    patient=crop('02_patient_twin.png',(265,360,1490,1210),'patient_crop.png')
    explain=crop('03_explainability.png',(265,540,1490,1150),'explain_crop.png')
    scenario=crop('04_what_if.png',(600,420,1490,980),'scenario_crop.png')
    care=crop('06_alerts_and_interventions.png',(265,350,1490,1000),'care_crop.png')
    # 1 cover: all native editable shapes/text.
    s=new_slide(prs,'','Happiest Health · Digital Twin Challenge 2026',1,dark=True)
    label(s,.62,1.4,3.55,'BODY × ENVIRONMENT DIGITAL TWIN',True)
    text(s,.62,2.08,9,1.0,'PranaGuard',59,WHITE,True)
    text(s,.66,3.34,8.0,1.05,'A connected respiratory twin.\nAn honest forecast of what comes next.',27,'DDEBD2')
    text(s,.68,4.85,8.0,.65,'COPD exacerbation forecasting · 24 / 48 / 72 hours',17,'A9C8AE')
    text(s,.68,5.92,8.4,.67,f'{team}\n{institution}',13,'9FBCA8')
    for i,h in enumerate([24,48,72]):
        y=2.15+i*1.1;rect(s,10.08,y,2.38,.8,fill='285346',line='3E6954');text(s,10.33,y+.14,1.92,.43,f'{h} hours',24,LIME,True)
    text(s,10.12,5.78,2.26,.6,'All patients synthetic.\nNot clinically validated.',11,'B7D0B4')
    # 2 clinical problem + Indian context.
    s=new_slide(prs,'Respiratory health is a moving picture.','01 / Problem & Indian healthcare use case',2,'A clinician needs changing patient context, not another disconnected measurement.')
    callout(s,.58,2.3,5.8,2.05,'The clinical question','Could a personalized combination of history, body signals, symptoms and environment support an earlier review conversation?',fill=WHITE)
    callout(s,.58,4.62,5.8,1.75,'The prototype boundary','This build explores the workflow. It does not demonstrate fewer admissions, clinical warning lead time or treatment benefit.',fill='F8EFE5',accent=CORAL)
    bullets(s,6.95,2.38,5.7,['Synthetic COPD outpatient cohort','Clinician review, acknowledgement and action logging','Pune · Mumbai · Delhi · Bengaluru ambient examples','No clinical participants or hospital partnership claimed'],18,.95)
    # 3 outputs.
    s=new_slide(prs,'Three horizons. Three actual models.','02 / Target outcome',3,'New synthetic onset in (t, t+h] — not diagnosis, triage or validated lead time.')
    for i,h in enumerate([24,48,72]):
        x=.6+i*4.22;rect(s,x,2.42,3.98,2.65,WHITE,LINE);text(s,x+.24,2.7,3.5,.8,str(h),53,TEAL,True);text(s,x+.28,3.66,3.42,.5,'HOUR HORIZON',14,MUTED,True);text(s,x+.28,4.3,3.42,.44,'Independent fitted + calibrated estimate',13,MUTED)
    text(s,.65,5.65,12.0,.8,'Active episodes, initial72h calibration and incomplete future follow-up are excluded from evaluation. Independent horizons can be non-monotonic.',18,MUTED)
    # 4 solution.
    s=new_slide(prs,'One clinician workspace, connected end to end.','03 / PranaGuard solution',4)
    picture_fit(s,overview,.6,2.15,8.0,4.48)
    bullets(s,9.02,2.22,3.7,['Cohort risk & source freshness','Patient-specific baseline and history','Explanations and isolated scenarios','Persisted alerts and review trail'],18,1.02)
    text(s,.63,6.72,11.9,.18,'Actual running React frontend screenshot. Values are API-fed synthetic model estimates, not a mockup.',9,MUTED)
    # 5 twin definition.
    s=new_slide(prs,'A stateful twin — not a static prediction page.','04 / Why it qualifies as a research digital twin',5)
    labels=[('01','Persistent identity','Patient-specific state and immutable SQLite versions.'),('02','Personal calibration','First72h baseline precedes fitted and replayed predictions.'),('03','Hourly updates','Observed body/environment context advances deterministically.'),('04','Future exploration','Three fitted risk models plus separate conditional branches.'),('05','Auditable review','Saved explanations-by-state, alerts, actions and scenarios.')]
    for i,(idx,title,body) in enumerate(labels):
        y=2.13+i*.88;label(s,.65,y,.55,idx);text(s,1.5,y-.03,3.1,.48,title,20,INK,True);text(s,4.72,y,7.65,.55,body,17,MUTED)
    # 6 diagram native editable scaled.
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG);architecture(s,scale=.82,ox=.1,oy=0,header=True)
    # 7 data provenance.
    s=new_slide(prs,'Know what is real. Know what is synthetic.','06 / Data provenance',7)
    callout(s,.62,2.15,6.0,3.18,'SYNTHETIC · clinical/body/outcomes','120 patients × 60 days = 172,800 hourly records\n263 stochastic artificial onsets\nSeed2026 · no real people or hospital records\nGenerator is not medically calibrated',fill=WHITE)
    callout(s,6.9,2.15,5.8,3.18,'REAL CACHED · ambient model products','8 successfully retrieved public API responses\n4 cities × air + weather; 1,440 hours/cache\n2026-08-01 to 2026-09-29 UTC\nOpen-Meteo / CAMS · CC BY 4.0',fill='E7F0DF')
    text(s,.67,5.72,12,.84,'Ambient grids are not personal exposure. Historical products are not live or verified as-issued forecasts. URLs, retrieval times, licenses and SHA-256 checksums are versioned.',18,MUTED)
    # 8 methodology.
    s=new_slide(prs,'The split comes before the score.','07 / AI & ML methodology',8,'No patient identity, future events, future exposure or full-course baseline averages enter the models.')
    for i,(name,count,body) in enumerate([('TRAIN',72,'14,605 sampled rows\nImputer/scaler fitted here'),('CALIBRATION',24,'1,377 sampled rows\nSeparate Platt fits here'),('TEST',24,'4,699 hourly windows\nLater dates; untouched for tuning')]):
        x=.62+i*4.22;rect(s,x,2.34,3.98,2.3,WHITE,LINE);text(s,x+.22,2.58,3.5,.35,name,12,TEAL,True);text(s,x+.22,3.05,3.5,.7,f'{count} patients',28,INK,True);text(s,x+.22,3.92,3.5,.55,body,14,MUTED)
    bullets(s,.73,5.0,12,['Patient-disjoint and calendar-time-disjoint, with ≥72h purges.','Fixed hyperparameters; LR baseline + 4 LightGBM input groups × 3 horizons.','Causal trailing features; first72h warm-up; no test-selected threshold or tuning.'],16,.56)
    # 9 explanations.
    s=new_slide(prs,'Explain the model — and its failure modes.','08 / Actual explainability',9)
    picture_fit(s,explain,.6,2.3,8.0,4.02)
    bullets(s,8.98,2.25,3.7,['Native exact TreeSHAP','Base + contributions = raw log-odds','Patient-relative deviations and driver deltas'],17,1.03)
    text(s,.65,6.42,12,.4,'Critical caveat: 72h Platt slope is negative. Raw contribution direction is not calibrated-risk direction. Association ≠ causation.',13,CORAL,True)
    # 10 scenario.
    s=new_slide(prs,'What changes when assumptions change?','09 / Conditional what-if studio',10,'A model sensitivity branch, not a validated treatment or physiological response simulator.')
    bullets(s,.65,2.4,4.1,['PM2.5 −80% to +100%','Adherence feature0–1','Activity0.5–1.5× steps','Observed physiology held constant','Original twin never overwritten'],17,.75)
    picture_fit(s,scenario,5.1,2.25,7.6,4.52)
    text(s,.68,6.4,4.05,.44,'Counterintuitive results are possible. No proven health benefit is implied.',12,CORAL,True)
    # 11 actual workflow.
    s=new_slide(prs,'From review to an auditable action.','10 / Actual frontend workflow',11,'Open a twin → inspect evidence → review an alert → document an action → reload persisted state.')
    picture_fit(s,patient,.62,2.3,5.4,4.33);picture_fit(s,care,6.3,2.3,6.4,4.33)
    text(s,.68,6.72,5.3,.18,'Actual patient twin, personal baseline and versioned risks',9,MUTED);text(s,6.36,6.72,6.2,.18,'Actual model-generated false-positive alert and persisted review note',9,MUTED)
    # 12 executed performance.
    s=new_slide(prs,'The engineering works. The model needs work.','11 / Actual executed evaluation',12,'Synthetic held-out results only. Poor72h discrimination and zero sensitivity at the fixed.35 cutoff.')
    metricrows=[]
    for h in [24,48,72]:
        for alg in ['LightGBM','LogisticRegression']:
            m=next(z for z in r['results'] if z['algorithm']==alg and z['ablation']=='personalized' and z['horizon']==h)
            metricrows.append([f'{h}h · '+('LightGBM' if alg=='LightGBM' else 'Logistic Reg.'),f'{m["auroc"]:.4f}',f'{m["auprc"]:.4f}',f'{m["brier"]:.4f}'])
    add_table(s,.62,2.3,7.6,3.9,['MODEL / HORIZON','AUROC','AUPRC','BRIER'],metricrows,15,widths=[3.25,1.45,1.45,1.45])
    xy=XyChartData();roc=xy.add_series('72h ROC (primary)')
    for point in r['diagnostics']['72']['roc']:roc.add_data_point(point['x'],point['y'])
    base=xy.add_series('Chance');base.add_data_point(0,0);base.add_data_point(1,1)
    chart=s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,Inches(8.62),Inches(2.5),Inches(4.05),Inches(3.33),xy).chart
    chart.has_legend=True;chart.legend.position=XL_LEGEND_POSITION.BOTTOM;chart.legend.font.name=FONT;chart.legend.font.size=Pt(9)
    for ax in [chart.category_axis,chart.value_axis]:ax.minimum_scale=0;ax.maximum_scale=1;ax.tick_labels.font.size=Pt(9);ax.tick_labels.font.name=FONT
    chart.series[0].format.line.color.rgb=rgb(TEAL);chart.series[0].format.line.width=Pt(2);chart.series[1].format.line.color.rgb=rgb('B8C6B4')
    text(s,8.82,6.08,3.65,.54,'72h prevalence:0.0773\n72h AUPRC:0.0593 (below baseline)',13,CORAL,True)
    text(s,.68,6.5,7.45,.3,'Test:24 patients,4,699 correlated hours.150 patient-cluster bootstrap resamples in JSON.',10,MUTED)
    # 13 ablations.
    s=new_slide(prs,'More features did not establish better forecasts.','12 / Actual ablation outcomes',13,'Same leakage-aware split, fixed model choices and separate calibration for every experiment.')
    abl=[z for z in r['results'] if z['algorithm']=='LightGBM' and z['horizon']==72]
    add_table(s,.62,2.3,7.6,3.12,['72H INPUT GROUP','AUROC','AUPRC','BRIER'],[[m['ablation'].replace('_',' ').title(),f'{m["auroc"]:.4f}',f'{m["auprc"]:.4f}',f'{m["brier"]:.4f}'] for m in abl],16,widths=[3.25,1.45,1.45,1.45])
    callout(s,8.62,2.3,4.05,3.12,'No improvement claim','Full environment72h is worse than body streams. Personalization does not establish benefit. Calibration instability is retained, not concealed.',fill='F8EEE5',accent=CORAL)
    text(s,.68,5.85,12,1.0,'Only5 held-out test-window onsets.0 detected at72h ≥35%; median lead time is undefined. All primary horizons have zero sensitivity at that cutoff. Small synthetic samples and shared exposure create major uncertainty.',19,CORAL)
    # 14 demo.
    s=new_slide(prs,'A reproducible demo, including a false positive.','13 / Deterministic demonstration scenario',14)
    steps=[('RESET','1039','24h persisted history · focus PG-0071'),('ADVANCE','+12h','Replay to1051 · state version37'),('REVIEW','35.51%','Actual calibrated72h risk crosses.35'),('DISCLOSE','label = 0','Retrospective score-selected false positive')]
    for i,(cap,big,body) in enumerate(steps):
        x=.63+i*3.18;rect(s,x,2.4,3.0,2.42,WHITE,LINE);text(s,x+.2,2.64,2.6,.3,cap,11,TEAL,True);text(s,x+.2,3.13,2.6,.72,big,32,INK,True);text(s,x+.2,4.0,2.6,.56,body,13,MUTED)
    callout(s,.66,5.18,12.0,1.37,'Show the full workflow','Play/pause → reset → +12h → explanation → acknowledge/log/resolve → scenario → reload → actual validation.',fill='E9F2E1')
    # 15 safety.
    s=new_slide(prs,'Clinical claims require clinical evidence.','14 / Safety, limitations & future validation',15)
    bullets(s,.7,2.2,6.0,['Synthetic generator and labels are not clinically calibrated.','72h model/calibration are unstable; no useful warning is established.','Ambient grids ≠ personal dose; scenarios are non-causal.','Local demo token ≠ production identity, privacy or governance.'],18,.98)
    callout(s,7.2,2.23,5.4,3.95,'Next, before real-world use','Ethics/consent + adjudicated real outcomes\nNew pre-registered development data\nExternal prospective patient/time/site validation\nCalibration, fairness and uncertainty assessment\nSecure ingestion and clinician safety governance',fill=WHITE)
    # 16 closing.
    s=new_slide(prs,'Build the workflow. Measure honestly.','15 / Conclusion & submission gate',16,dark=True)
    text(s,.66,2.15,8.3,1.03,'Persistent state. Real inference.\nTransparent limitations.',35,WHITE,True)
    text(s,.7,3.64,8.4,.83,'120 synthetic twins · 60 days · 3 horizons\nReproducible experiments and a working clinician workspace.',19,'BBD1B9')
    for i,item in enumerate(['Team / institution details: PENDING','Public GitHub + signed-out file checks: PENDING','Actual20min Unlisted YouTube upload: PENDING']):text(s,.72,5.1+i*.44,11.8,.35,item,14,LIME)
    text(s,.72,6.62,11.9,.28,'NOT FULLY SUBMISSION-READY until the mandatory external items are genuinely complete.',12,'D5BDA1',True)
    prs.save(PDIR/'PranaGuard_Digital_Twin_Challenge_2026.pptx')
    print('Created editable16-slide deck and separate editable architecture.')

def export_pdfs():
    for folder,name in [(ADIR,'PranaGuard_Architecture.pptx'),(PDIR,'PranaGuard_Digital_Twin_Challenge_2026.pptx')]:
        cmd=['libreoffice','-env:UserInstallation=file:///home/user/.cache/pranaguard-libreoffice','--headless','--convert-to','pdf','--outdir',str(folder),str(folder/name)]
        subprocess.run(cmd,check=True,timeout=120)
        pdf=folder/Path(name).with_suffix('.pdf');assert pdf.exists() and pdf.stat().st_size>1000
    document=fitz.open(ADIR/'PranaGuard_Architecture.pdf');(ADIR/'PranaGuard_Architecture.svg').write_text(document[0].get_svg_image(text_as_path=True));document.close()

def technical_pdf():
    regular=next(Path('/usr/share/fonts/truetype').rglob('LiberationSans-Regular.ttf'));bold=next(Path('/usr/share/fonts/truetype').rglob('LiberationSans-Bold.ttf'))
    pdfmetrics.registerFont(TTFont('PGSans',str(regular)));pdfmetrics.registerFont(TTFont('PGSansBold',str(bold)))
    styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='PGBody',fontName='PGSans',fontSize=10,leading=15,spaceAfter=9,textColor=colors.HexColor('#25483b')));styles.add(ParagraphStyle(name='PGTitle',fontName='PGSansBold',fontSize=23,leading=29,spaceAfter=20,textColor=colors.HexColor('#163e35')));styles.add(ParagraphStyle(name='PGHeading',fontName='PGSansBold',fontSize=14,leading=19,spaceBefore=17,spaceAfter=9,textColor=colors.HexColor('#287e68')));styles.add(ParagraphStyle(name='PGSmall',fontName='PGSans',fontSize=8,leading=11,spaceAfter=3))
    lines=(DOCS/'TECHNICAL_REPORT.md').read_text().splitlines();story=[];i=0
    def para(t,style):return Paragraph(escape(t).replace('`',''),styles[style])
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                vals=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch('[-: ]+',v) for v in vals):rows.append([para(v,'PGSmall') for v in vals])
                i+=1
            n=len(rows[0]);widths=[222]+[(483-222)/(n-1)]*(n-1)
            t=Table(rows,colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e6efdf')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f5f8f0')]),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,-1),.3,colors.HexColor('#dce6d5')),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]));story.extend([t,Spacer(1,12)]);continue
        if line.startswith('# '):story.append(para(line[2:],'PGTitle'))
        elif line.startswith('## '):story.append(para(line[3:],'PGHeading'))
        elif line.startswith('### '):story.append(para(line[4:],'PGHeading'))
        else:
            group=[line]
            while i+1<len(lines) and lines[i+1].strip() and not lines[i+1].startswith(('#','|')):i+=1;group.append(lines[i].strip())
            story.append(para(' '.join(group),'PGBody'))
        i+=1
    def on_page(canvas,doc):
        canvas.setStrokeColor(colors.HexColor('#dce6d5'));canvas.line(50,45,545,45);canvas.setFont('PGSans',8);canvas.setFillColor(colors.HexColor('#718569'));canvas.drawString(50,31,'PRANAGUARD · SYNTHETIC RESEARCH PROTOTYPE · NOT CLINICALLY VALIDATED');canvas.drawRightString(545,31,str(doc.page))
    doc=SimpleDocTemplate(str(DOCS/'PranaGuard_Technical_Report.pdf'),pagesize=(595.28,841.89),rightMargin=55,leftMargin=55,topMargin=50,bottomMargin=64,title='PranaGuard technical report',author='PranaGuard contributors (team details pending)')
    doc.build(story,onFirstPage=on_page,onLaterPages=on_page)

def build():
    for p in [PDIR,ADIR,ASSETS]:p.mkdir(parents=True,exist_ok=True)
    build_architecture();build_deck();export_pdfs();technical_pdf()
if __name__=='__main__':build()
