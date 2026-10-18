"""Inspect actual Office/PDF content, page/shape bounds, text preservation and image decodes."""
import json,re,unicodedata,zipfile
from datetime import datetime,timezone
from pathlib import Path
from PIL import Image,ImageOps,ImageDraw
from pptx import Presentation
import pymupdf
from backend.app.config import ROOT

def norm(text):return ''.join(c for c in unicodedata.normalize('NFKC',text).casefold() if c.isalnum())
def paragraphs(shape):
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:yield p.text
    if shape.has_table:
        for row in shape.table.rows:
            for c in row.cells:
                for p in c.text_frame.paragraphs:yield p.text

def verify():
    reports=[];render=Path('/home/user/.cache/pranaguard-artifacts');render.mkdir(parents=True,exist_ok=True)
    pairs=[('docs/architecture/PranaGuard_Architecture',1),('docs/presentation/PranaGuard_Digital_Twin_Challenge_2026',16)]
    thumbs=[]
    for stem,count in pairs:
        pptx=ROOT/(stem+'.pptx');pdf=ROOT/(stem+'.pdf')
        with zipfile.ZipFile(pptx) as z:assert z.testzip() is None
        prs=Presentation(pptx);doc=pymupdf.open(pdf);assert len(prs.slides)==len(doc)==count
        outside=[];missing=[];span_outside=[];font_min=100
        for i,slide in enumerate(prs.slides):
            page=doc[i];pdf_text=norm(page.get_text())
            for shape in slide.shapes:
                if shape.left<-10000 or shape.top<-10000 or shape.left+shape.width>prs.slide_width+10000 or shape.top+shape.height>prs.slide_height+10000:outside.append({'slide':i+1,'shape':shape.name})
                for expected in paragraphs(shape):
                    n=norm(expected)
                    if len(n)>=12 and n not in pdf_text:missing.append({'slide':i+1,'text':expected})
            for block in page.get_text('dict')['blocks']:
                if 'lines' not in block:continue
                for line in block['lines']:
                    for span in line['spans']:
                        bbox=span['bbox'];font_min=min(font_min,span['size'])
                        if bbox[0]<-.5 or bbox[1]<-.5 or bbox[2]>page.rect.width+.5 or bbox[3]>page.rect.height+.5:span_outside.append({'page':i+1,'text':span['text'],'bbox':bbox})
            pix=page.get_pixmap(matrix=pymupdf.Matrix(1.2,1.2),alpha=False);img_path=render/(Path(stem).name+f'_{i+1:02d}.png');pix.save(img_path)
            if count==16:
                im=Image.open(img_path);im.thumbnail((480,270));card=Image.new('RGB',(500,300),'#e5ece4');card.paste(im,((500-im.width)//2,10));ImageDraw.Draw(card).text((12,282),f'SLIDE {i+1:02d}',fill='#173c32');thumbs.append(card)
        reports.append({'pptx':str(pptx.relative_to(ROOT)),'pdf':str(pdf.relative_to(ROOT)),'slides':count,'pages':len(doc),'ooxml_zip_integrity':True,'native_editable_text_verified':True,'minimum_rendered_font_pt':font_min,'shapes_outside_slide':outside,'text_missing_from_pdf_export':missing,'pdf_spans_outside_page':span_outside,'status':'PASS' if not outside and not missing and not span_outside else 'FAIL'})
        doc.close()
    technical=pymupdf.open(ROOT/'docs/PranaGuard_Technical_Report.pdf');text=''.join(p.get_text() for p in technical)
    assert len(technical)>=5 and '0.3951' in text and '0.6253' in text and 'clinical' in text.lower()
    reports.append({'pdf':'docs/PranaGuard_Technical_Report.pdf','pages':len(technical),'metric_text_verified':True,'status':'PASS'})
    decoded=[]
    for f in (ROOT/'docs/screenshots').glob('*.png'):
        im=Image.open(f);im.verify();decoded.append(str(f.relative_to(ROOT)))
    assert len(decoded)>=8
    sheet=Image.new('RGB',(2000,1200),'#ccd8ce')
    for i,im in enumerate(thumbs):sheet.paste(im,((i%4)*500,(i//4)*300))
    sheet.save(render/'deck-review.png')
    data={'created_at':datetime.now(timezone.utc).isoformat(),'reports':reports,'screenshots_decoded':decoded,'all_pass':all(r['status']=='PASS' for r in reports),'visual_review':'Render review required in addition to these automated checks; see QA notes.','notes':['PPTX parsed with python-pptx and ZIP integrity checked.','PDFs are actual LibreOffice exports, not separately invented slide content.','Every substantial native paragraph/table cell checked against the corresponding exported PDF text.','Bounds check detects page overflow, not every possible semantic layout issue.']}
    (ROOT/'docs/verification/artifact-verification.json').write_text(json.dumps(data,indent=2))
    print(json.dumps(data,indent=2));assert data['all_pass'],'Fix missing/clipped content before claiming completion.'
if __name__=='__main__':verify()
